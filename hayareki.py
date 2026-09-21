#!/usr/bin/env python3
"""はやれき — 栽培記録の検算と、JA提出様式への書き戻し

使い方:
  python hayareki.py check  --db 記録.xlsx
  python hayareki.py export --db 記録.xlsx --template 様式.xlsx --out 提出用.xlsx

--db には次のどちらでも渡せる（自動判別）:
  - はやれき入力アプリのスプレッドシートを xlsx でダウンロードしたもの
  - AppSheet 版の DB スプレッドシートを xlsx でダウンロードしたもの
"""
import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


# ---------------------------------------------------------------- 共通
def key(s):
    """表記ゆれ吸収用のキー（全角英数→半角、空白除去）"""
    if s is None:
        return ""
    return unicodedata.normalize("NFKC", str(s)).replace(" ", "").replace("\u3000", "").strip()


def num(v, default=None):
    if v is None or v == "":
        return default
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return default


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if v is None or v == "":
        return None
    s = str(v).strip().split(" ")[0].replace("-", "/")
    for fmt in ("%Y/%m/%d", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def read_tables(path):
    """各シートの先頭行を見出しとして dict の行リストを返す"""
    wb = load_workbook(path, data_only=True, read_only=True)
    tables = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows[:5]) if r and sum(c is not None for c in r) >= 2), None)
        if hi is None:
            continue
        head = [key(h) for h in rows[hi]]
        data = []
        for r in rows[hi + 1:]:
            if not r or all(c is None or c == "" for c in r):
                continue
            data.append({head[i]: r[i] for i in range(min(len(head), len(r))) if head[i]})
        tables[ws.title] = (head, data)
    return tables


def find(tables, *need, forbid=()):
    for name, (head, data) in tables.items():
        hs = set(head)
        if all(any(n in h for h in hs) for n in need) and not any(f in hs for f in forbid):
            return data
    return None


# ------------------------------------------------------------- データ形
# 正規化後の記録:
#   {event, date, plot, cat, item, qty, unit, water, flower, harvest, note}
# 防除の qty / water はそのタンク全体の量（区画ごとに行が分かれても同じ値）


def expand_plots(s, cfg):
    alias = {key(k): v for k, v in cfg.get("plots_alias", {}).items()}
    out = []
    for p in re.split(r"[,、，]", str(s or "")):
        p = key(p)
        if not p:
            continue
        out.extend(key(x) for x in alias.get(p, [p]))
    return out


def load_hayareki(tables):
    rec = find(tables, "event_id", "分類", "項目")
    plots = find(tables, "区画名", "ハウス数")
    ferts = find(tables, "資材名", "化学N%")
    pests = find(tables, "薬剤名", "使用回数上限")
    if rec is None:
        return None
    M = {"plots": {}, "ferts": {}, "pests": {}, "disp": {}}
    for r in plots or []:
        M["plots"][key(r["区画名"])] = {"area": num(r.get("面積a"), 0), "houses": num(r.get("ハウス数"))}
        M["disp"][key(r["区画名"])] = str(r["区画名"])
    for r in ferts or []:
        M["ferts"][key(r["資材名"])] = {
            "n": num(r.get("全N%"), 0), "cn": num(r.get("化学N%"), 0),
            "p": num(r.get("リン酸%"), 0), "k": num(r.get("カリ%"), 0)}
    for r in pests or []:
        M["pests"][key(r["薬剤名"])] = {
            "cls": r.get("分類") or "", "max": num(r.get("使用回数上限"), 0),
            "count": num(r.get("カウント"), None), "method": r.get("計算方法") or "希釈",
            "min": num(r.get("最小倍率")), "maxr": num(r.get("最大倍率")), "unit": r.get("単位") or "ml"}
    R = []
    for r in rec:
        d = to_date(r.get("日付"))
        if not d:
            continue
        R.append({"event": str(r.get("event_id")), "date": d, "plot": key(r.get("区画")),
                  "cat": r.get("分類"), "item": key(r.get("項目")), "qty": num(r.get("使用量")),
                  "unit": r.get("単位") or "", "water": num(r.get("水量L")),
                  "flower": num(r.get("開花段")), "harvest": num(r.get("収穫段")), "note": r.get("メモ") or ""})
    return R, M


def load_appsheet(tables, cfg):
    fert = find(tables, "記録ID", "資材名", "施用量")
    pest = find(tables, "記録ID", "散布液量(L)", "農薬1")
    work = find(tables, "記録ID", "作業内容", "開花段数")
    mast = find(tables, "区画リスト", "資材リスト")
    pmas = find(tables, "農薬名", "使用回数上限")
    if fert is None and pest is None:
        return None
    ap = cfg.get("fertilizer_apportion") or {}
    aph = ap.get("area_per_house", 2)
    M = {"plots": {}, "ferts": {}, "pests": {}, "disp": {}}
    for r in mast or []:
        if r.get("区画リスト"):
            a = num(r.get("面積(a)"), 0)
            M["plots"][key(r["区画リスト"])] = {"area": a, "houses": a / aph if aph else None}
            M["disp"][key(r["区画リスト"])] = str(r["区画リスト"])
        if r.get("資材リスト"):
            pct = lambda x: (num(r.get(x), 0) or 0) * 100
            M["ferts"][key(r["資材リスト"])] = {"n": pct("全チッソ"), "cn": pct("化学N"),
                                               "p": pct("リン酸"), "k": pct("カリ")}
    rule = cfg.get("count_rule", {})
    for r in pmas or []:
        if not r.get("農薬名"):
            continue
        mx = num(r.get("使用回数上限"), 0)
        cls = r.get("分類") or ""
        excluded = cls in rule.get("exclude_classes", []) or (rule.get("exclude_if_unlimited") and mx == 0)
        bai = num(r.get("希釈倍率"))
        M["pests"][key(r["農薬名"])] = {"cls": cls, "max": mx, "count": 0 if excluded else 1,
                                       "method": "希釈" if bai else "直", "min": bai, "maxr": None,
                                       "unit": r.get("使用単位") or "ml"}
    R = []
    for r in fert or []:
        d = to_date(r.get("日付"))
        if not d:
            continue
        for p in expand_plots(r.get("区画"), cfg):
            R.append({"event": str(r.get("記録ID")), "date": d, "plot": p, "cat": "施肥",
                      "item": key(r.get("資材名")), "qty": num(r.get("施用量(Kg)")), "unit": "kg",
                      "water": None, "flower": None, "harvest": None, "note": r.get("備考") or ""})
    for r in pest or []:
        d = to_date(r.get("日付"))
        if not d:
            continue
        water = num(r.get("散布液量(L)"))
        for i in (1, 2, 3):
            item = key(r.get(f"農薬{i}"))
            if not item:
                continue
            unit = M["pests"].get(item, {}).get("unit", "ml")
            for p in expand_plots(r.get("区画"), cfg):
                R.append({"event": f'{r.get("記録ID")}-{i}', "date": d, "plot": p, "cat": "防除",
                          "item": item, "qty": num(r.get(f"使用量{i}")), "unit": unit, "water": water,
                          "flower": None, "harvest": None, "note": r.get("備考") or ""})
    for r in work or []:
        d = to_date(r.get("日付"))
        if not d:
            continue
        for p in expand_plots(r.get("区画"), cfg):
            R.append({"event": str(r.get("記録ID")), "date": d, "plot": p, "cat": "作業",
                      "item": key(r.get("作業内容")), "qty": None, "unit": "", "water": None,
                      "flower": num(r.get("開花段数")), "harvest": num(r.get("収穫段数")),
                      "note": r.get("備考") or ""})
    return R, M


def load(path, cfg):
    t = read_tables(path)
    got = load_hayareki(t)
    src = "はやれき入力アプリ"
    if got is None:
        got = load_appsheet(t, cfg)
        src = "AppSheet DB"
    if got is None:
        sys.exit("記録のシートが見つかりません。入力アプリかAppSheet DBをxlsxで書き出したものを指定してください。")
    R, M = got
    skip = {key(s) for s in cfg.get("skip_items", [])}
    R = [r for r in R if r["item"] not in skip]
    return R, M, src


# ---------------------------------------------------------------- 検算
def fert_ratio(plot, M, cfg):
    ap = cfg.get("fertilizer_apportion") or {}
    basis = ap.get("basis_houses")
    h = M["plots"].get(plot, {}).get("houses")
    return (h / basis) if (basis and h) else 1.0


def kg(qty, unit):
    return (qty or 0) * (0.001 if key(unit) in ("g", "ml") else 1)


def pest_events(R):
    """(event, item) 単位にまとめる。防除の量は区画で重複させない"""
    ev = {}
    for r in R:
        if r["cat"] != "防除":
            continue
        k = (r["event"], r["item"])
        e = ev.setdefault(k, {"date": r["date"], "item": r["item"], "qty": r["qty"],
                              "water": r["water"], "unit": r["unit"], "plots": []})
        if r["plot"] not in e["plots"]:
            e["plots"].append(r["plot"])
    return list(ev.values())


def check(R, M, cfg):
    warn = []
    # 希釈倍率
    ratio_rows = []
    for e in sorted(pest_events(R), key=lambda x: (x["date"], x["item"])):
        m = M["pests"].get(e["item"])
        if not m:
            warn.append(f'農薬マスターにない薬剤: {e["item"]}')
            continue
        if m["method"] != "希釈" or not e["qty"] or not e["water"]:
            continue
        r = e["water"] * 1000 / e["qty"]
        st = "○"
        if m["min"] and r < m["min"]:
            st = f'濃すぎ（基準{m["min"]:g}倍以上）'
        elif m["maxr"] and r > m["maxr"]:
            st = f'薄すぎ（基準{m["maxr"]:g}倍以下）'
        ratio_rows.append((e["date"], e["item"], "・".join(M.get("disp", {}).get(x, x) for x in e["plots"]), e["qty"], e["water"], r, st))
    # 使用回数（区画×薬剤で散布回数を数える）
    uses = defaultdict(set)
    for r in R:
        if r["cat"] == "防除":
            uses[(r["plot"], r["item"])].add(r["event"])
    count_rows = []
    total = defaultdict(int)
    for (p, it), evs in sorted(uses.items()):
        m = M["pests"].get(it, {})
        n = len(evs)
        mx = m.get("max") or 0
        st = "制限なし" if mx == 0 else ("超過" if n > mx else ("上限" if n == mx else "○"))
        count_rows.append((p, it, n, mx, st))
        if m.get("count", 1):
            total[p] += n
    # 化学窒素・カリ
    n_plot = defaultdict(float)
    nk = defaultdict(lambda: [0.0, 0.0])
    for r in R:
        if r["cat"] != "施肥":
            continue
        f = M["ferts"].get(r["item"])
        if not f:
            warn.append(f'肥料マスターにない資材: {r["item"]}')
            continue
        q = kg(r["qty"], r["unit"]) * fert_ratio(r["plot"], M, cfg)
        n_plot[r["plot"]] += q * f["cn"] / 100
        mm = nk[(r["plot"], r["date"].month)]
        mm[0] += q * f["cn"] / 100
        mm[1] += q * f["k"] / 100
    return {"ratio": ratio_rows, "count": count_rows, "total": total,
            "n": n_plot, "nk": nk, "warn": sorted(set(warn))}


def print_check(res, M, cfg, src):
    D = lambda p: M.get("disp", {}).get(p, p)
    lim = cfg.get("n_limit_per_a", 4)
    cmax = cfg.get("count_rule", {}).get("count_limit_total", 36)
    print(f"読み込み元: {src}\n")
    print("■ 希釈倍率")
    for d, it, pl, q, w, r, st in res["ratio"]:
        mark = "  " if st == "○" else "⚠ "
        print(f"{mark}{d:%m/%d} {it:<16} {q:g}/{w:g}L = {r:,.1f}倍  {st}  [{pl}]")
    print("\n■ 使用回数（上限に達したもの・超えたもの）")
    hit = [x for x in res["count"] if x[4] in ("上限", "超過")]
    for p, it, n, mx, st in hit:
        print(f"  {D(p):<8} {it:<16} {n}/{mx:g}回  {st}")
    if not hit:
        print("  なし")
    print(f"\n■ 区画別カウント（記録分のみ・育苗期など記録外は含まない／上限{cmax}）")
    for p in M["plots"]:
        if res["total"].get(p):
            print(f"  {D(p):<8} {res['total'][p]}回  残り{cmax - res['total'][p]}")
    print(f"\n■ 化学窒素（上限 {lim:g}kg/a）")
    for p, v in M["plots"].items():
        if p in res["n"]:
            L = v["area"] * lim
            print(f"  {D(p):<8} {res['n'][p]:.2f}kg / {L:g}kg  残り{L - res['n'][p]:.2f}")
    print("\n■ 月別 K:N 比（カリ÷化学窒素）")
    months = sorted({m for (_, m) in res["nk"]})
    print("  " + " " * 8 + "".join(f"{m:>7}月" for m in months))
    for p in M["plots"]:
        cells = []
        for m in months:
            n, k = res["nk"].get((p, m), (0, 0))
            cells.append(f"{k / n:>8.2f}" if n > 0 else f"{'-':>8}")
        if any(c.strip() != "-" for c in cells):
            print(f"  {D(p):<8}" + "".join(cells))
    if res["warn"]:
        print("\n■ 注意")
        for w in res["warn"]:
            print("  " + w)


# ------------------------------------------------------------ 様式出力
def build_listmap(wb, cfg):
    """様式内の肥料・農薬リストから、表記ゆれを吸収して正式表記を引く辞書を作る"""
    lm = {}
    for k in ("fertilizer_list", "pesticide_list"):
        L = cfg["ja_form"].get(k)
        if not L or L["sheet"] not in wb.sheetnames:
            continue
        ws = wb[L["sheet"]]
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, L["name_col"]).value
            if v and str(v).strip() not in ("―――", ""):
                lm.setdefault(key(v).casefold(), str(v).strip())
    return lm


def ja_name(item, cfg, lm, log=None):
    m = {key(k).casefold(): v for k, v in cfg.get("names_to_ja", {}).items()}
    name = m.get(item.casefold(), item)
    exact = lm.get(key(name).casefold())
    if exact:
        return exact
    if log is not None:
        log.append(f"注意: 様式のリストにない名前 → {name}（様式のリストに追記してください）")
    return name


def shift_formula(f, src_row, dst_row):
    return re.sub(rf"(\$?[A-Z]{{1,3}}\$?){src_row}(?!\d)", lambda m: f"{m.group(1)}{dst_row}", f)


def export(R, M, cfg, template, out):
    wb = load_workbook(template)
    J = cfg["ja_form"]
    log = []
    lm = build_listmap(wb, cfg)

    # 施肥: 月ブロックごとに資材×区画の合計を書く
    F = J["fertilizer"]
    ws = wb[F["sheet"]]
    fcols = {key(k): v for k, v in F["plot_cols"].items()}
    agg = defaultdict(float)
    for r in R:
        if r["cat"] == "施肥" and r["plot"] in fcols:
            agg[(r["date"].month, ja_name(r["item"], cfg, lm, log), r["plot"])] += \
                kg(r["qty"], r["unit"]) * fert_ratio(r["plot"], M, cfg)
    tr = F["formula_template_row"]
    for mon in sorted({k[0] for k in agg}):
        blk = F["month_blocks"].get(str(mon))
        if not blk:
            log.append(f"施肥: {mon}月のブロックが様式にありません（config の month_blocks を確認）")
            continue
        r0, r1 = blk
        have = {key(ws.cell(r, F["name_col"]).value): r for r in range(r0, r1 + 1) if ws.cell(r, F["name_col"]).value}
        free = [r for r in range(r0, r1 + 1) if not ws.cell(r, F["name_col"]).value]
        for mat in sorted({k[1] for k in agg if k[0] == mon}):
            row = have.get(key(mat))
            if row is None:
                if not free:
                    log.append(f"施肥: {mon}月のブロックに空き行がありません（{mat}）")
                    continue
                row = free.pop(0)
                ws.cell(row, F["name_col"], mat)
                for c in F.get("formula_cols", []):
                    f = ws.cell(tr, c).value
                    if isinstance(f, str) and f.startswith("="):
                        ws.cell(row, c, shift_formula(f, tr, row))
            for p, c in fcols.items():
                v = agg.get((mon, mat, p))
                if v:
                    ws.cell(row, c, round(v, 3))
            log.append(f"施肥 {mon}月 行{row}: {mat}")

    # 防除: 同日同剤を1行にまとめて追記（既に書かれている日×剤はとばす）
    P = J["pesticide"]
    ws = wb[P["sheet"]]
    pcols = {key(k): v for k, v in P["plot_cols"].items()}
    have = set()
    row = P["start_row"]
    while ws.cell(row, P["name_col"]).value:
        have.add((ws.cell(row, P["month_col"]).value, ws.cell(row, P["day_col"]).value,
                  key(ws.cell(row, P["name_col"]).value)))
        row += 1
    day = {}
    for e in pest_events(R):
        k = (e["date"], ja_name(e["item"], cfg, lm, log))
        d = day.setdefault(k, {"qty": 0, "water": 0, "plots": [], "item": e["item"], "unit": e["unit"]})
        d["qty"] += e["qty"] or 0
        d["water"] += e["water"] or 0
        for p in e["plots"]:
            if p not in d["plots"]:
                d["plots"].append(p)
    for (date, name), d in sorted(day.items()):
        if (date.month, date.day, key(name)) in have:
            continue
        m = M["pests"].get(d["item"], {})
        ws.cell(row, P["month_col"], date.month)
        ws.cell(row, P["day_col"], date.day)
        ws.cell(row, P["name_col"], name)
        ws.cell(row, P["qty_col"], d["qty"])
        ws.cell(row, P["unit_col"], m.get("unit") or d["unit"])
        ws.cell(row, P["water_col"], d["water"])
        counted = m.get("count", 1)
        if counted:
            for p in d["plots"]:
                if p in pcols:
                    ws.cell(row, pcols[p], 1)
        log.append(f"防除 行{row}: {date:%m/%d} {name} {d['qty']:g}/{d['water']:g}L"
                   f" 区画{len(d['plots'])}{'' if counted else '（カウント外）'}")
        row += 1

    wb.save(out)
    seen, uniq = set(), []
    for x in log:
        if x not in seen:
            seen.add(x); uniq.append(x)
    return uniq


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description="はやれき: 栽培記録の検算とJA様式出力")
    ap.add_argument("command", choices=["check", "export"])
    ap.add_argument("--db", required=True, help="記録スプレッドシートを xlsx で書き出したもの")
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--template", help="JA提出様式（前月の提出用でも可）")
    ap.add_argument("--out", help="出力ファイル名")
    a = ap.parse_args()

    cfg = json.loads(Path(a.config).read_text(encoding="utf-8"))
    R, M, src = load(a.db, cfg)
    if a.command == "check":
        print_check(check(R, M, cfg), M, cfg, src)
    else:
        if not a.template or not a.out:
            sys.exit("export には --template と --out が必要です")
        for line in export(R, M, cfg, a.template, a.out):
            print(line)
        print(f"\n書き出し: {a.out}（Excelで開くと計算式が再計算されます）")


if __name__ == "__main__":
    main()
